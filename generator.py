import html
import io
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF
import config

_REPLACE = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...", "\u00a0": " ",
}


def sanitize_text(text: str) -> str:
    """Remove typographic characters and markdown symbols."""
    for k, v in _REPLACE.items():
        text = text.replace(k, v)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"^#+\s*", "", text, flags=re.MULTILINE)
    return text.replace("`", "").strip()


def _is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    return bool(re.match(r"^(\d+\.|[IVX]+\.)\s", s)) or s.endswith(":") or (s.isupper() and len(s) > 3)


def split_terms(terms: str):
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def _latin(s: str) -> str:
    return s.encode("latin-1", "replace").decode("latin-1")


# ---------------- DOCX ----------------
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if config.LOGO_PATH.exists():
        doc.add_picture(str(config.LOGO_PATH), width=Inches(2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("LegalEase").bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.title())
    r.bold = True
    r.font.size = Pt(16)

    for line in text.split("\n"):
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.add_run(line.strip()).bold = _is_heading(line)

    items = split_terms(terms)
    if items:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = config.FOOTER_TEXT
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
class _PDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        if config.LOGO_PATH.exists():
            self.image(str(config.LOGO_PATH), x=(self.w - 40) / 2, y=8, w=40)
            self.set_y(30)
        else:
            self.set_font("Helvetica", "B", 16)
            self.cell(0, 10, "LegalEase", align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, _latin(self.doc_type.title()), align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"{config.FOOTER_TEXT}  Page {self.page_no()}", align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    pdf = _PDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    for line in text.split("\n"):
        if not line.strip():
            continue
        pdf.set_font("Helvetica", "B" if _is_heading(line) else "", 11)
        pdf.multi_cell(0, 6, _latin(line.strip()), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)
    items = split_terms(terms)
    if items:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 6, "Summary of Key Terms", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for t in items:
            pdf.multi_cell(0, 6, _latin("-  " + t), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ---------------- HTML preview ----------------
def format_html_preview(text: str) -> str:
    out = []
    for line in sanitize_text(text).split("\n"):
        if not line.strip():
            continue
        safe = html.escape(line.strip())
        out.append(f"<h4>{safe}</h4>" if _is_heading(line) else f"<p>{safe}</p>")
    return "".join(out)
